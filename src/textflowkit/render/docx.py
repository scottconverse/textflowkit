"""DOCX renderer.

Binary output, so it does not go through `render()` - see `render_bytes()`.

Renders the **finished** data model: speaker labels and translated text are both
included, which is why this format is built after those stages exist rather than
before.
"""

from __future__ import annotations

import io

from textflowkit.core.model import Transcript
from textflowkit.core.timeutil import srt_timestamp

try:
    from docx import Document
except ImportError as exc:  # pragma: no cover - optional extra
    raise ImportError(
        "DOCX export requires the 'export' extra. Install with: pip install 'textflowkit[export]'"
    ) from exc


def _hms(seconds: float) -> str:
    return srt_timestamp(seconds).split(",")[0]


def render_docx(transcript: Transcript, *, title: str = "Transcript") -> bytes:
    """Render a transcript as a .docx document."""
    doc = Document()
    doc.add_heading(title, level=1)

    meta = []
    if transcript.source:
        meta.append(f"Source: {transcript.source}")
    if transcript.language:
        meta.append(f"Language: {transcript.language}")
    if transcript.duration:
        meta.append(f"Duration: {_hms(transcript.duration)}")
    diar = (transcript.metadata or {}).get("diarization")
    if diar:
        speakers = ", ".join(diar.get("speakers") or []) or "none"
        meta.append(f"Speakers: {speakers}")
    trans = (transcript.metadata or {}).get("translation")
    if trans:
        meta.append(f"Translated to: {trans.get('target')}")

    for line in meta:
        doc.add_paragraph(line)
    doc.add_paragraph("")

    current_speaker: str | None = None
    for segment in transcript.segments:
        if segment.hidden:
            continue
        if segment.speaker and segment.speaker != current_speaker:
            current_speaker = segment.speaker
            doc.add_heading(current_speaker, level=2)

        para = doc.add_paragraph()
        para.add_run(f"[{_hms(segment.start)}] ").bold = True
        if segment.translated_text:
            para.add_run(segment.translated_text.strip())
            para.add_run(f"\n(original: {segment.text.strip()})").italic = True
        else:
            para.add_run(segment.text.strip())

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
